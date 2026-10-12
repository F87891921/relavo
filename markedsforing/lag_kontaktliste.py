"""Lager kontaktliste.xlsx fra offentlige registre (SSB og Enhetsregisteret).
Kjør med et miljø som har openpyxl: python3 lag_kontaktliste.py
Lista er på virksomhetsnivå. Personer fylles ut manuelt i kolonnene til høyre.
"""
import json, re, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def get(url, data=None):
    req = urllib.request.Request(url, data=json.dumps(data).encode() if data else None,
                                 headers={"Content-Type": "application/json", "Accept": "application/json"})
    return json.load(urllib.request.urlopen(req, timeout=90))

# ---------------------------------------------------------------- Kommuner
T = "https://data.ssb.no/api/v0/no/table/06913"
meta = get(T)
region = next(v for v in meta["variables"] if v["code"] == "Region")
tid = next(v for v in meta["variables"] if v["code"] == "Tid")
koder = [k for k in region["values"] if re.fullmatch(r"\d{4}", k)]
pop = {}
for ar in (tid["values"][-1], tid["values"][-2]):
    d = get(T, {"query": [{"code": "Region", "selection": {"filter": "item", "values": koder}},
                          {"code": "ContentsCode", "selection": {"filter": "item", "values": ["Folkemengde"]}},
                          {"code": "Tid", "selection": {"filter": "item", "values": [ar]}}],
                "response": {"format": "json-stat2"}})
    idx = d["dimension"]["Region"]["category"]["index"]; lab = d["dimension"]["Region"]["category"]["label"]
    pop = {k: d["value"][i] for k, i in idx.items() if d["value"][i]}
    if len(pop) > 300: break

kommuner_brreg, side = [], 0
while True:
    r = get(f"https://data.brreg.no/enhetsregisteret/api/enheter?organisasjonsform=KOMM&size=200&page={side}")
    kommuner_brreg += r.get("_embedded", {}).get("enheter", [])
    if side + 1 >= r["page"]["totalPages"]: break
    side += 1
by = {}
for e in kommuner_brreg:
    nr = (e.get("forretningsadresse") or e.get("postadresse") or {}).get("kommunenummer")
    if nr in pop: by.setdefault(nr, []).append(e)

def pris(n): return 3900 if n < 10000 else 7900 if n < 25000 else 12900 if n < 50000 else None
fmt = lambda x: f"{x:,}".replace(",", " ")
kommuner = []
for k, n in sorted(pop.items(), key=lambda kv: -kv[1]):
    es = by.get(k, [])
    e = next((x for x in es if x["navn"].upper().endswith("KOMMUNE")), es[0] if es else {})
    p = pris(n)
    seg = "A (25 000+)" if n >= 25000 else "B (10 000–25 000)" if n >= 10000 else "C (under 10 000)"
    kommuner.append([k, lab[k].split(" - ")[0], n, seg, fmt(p) if p else "Offert", fmt(p * 12) if p else "Offert",
                     e.get("organisasjonsnummer", ""), e.get("hjemmeside", ""), e.get("epostadresse", ""),
                     (e.get("telefon") or "").replace("\xa0", " ")])

# ---------------------------------------------------------------- Selskap
KODER = [("41", "Bygg: hovedentreprenør og utvikling"), ("42", "Anlegg"),
         ("43", "Bygg: spesialisert (typisk underentreprenør)"), ("81.21", "Renhold"), ("81.22", "Renhold (annet)")]
sel = {}
for kode, etikett in KODER:
    side = 0
    while True:
        r = get(f"https://data.brreg.no/enhetsregisteret/api/enheter?naeringskode={kode}&organisasjonsform=AS&fraAntallAnsatte=20&size=200&page={side}")
        for e in r.get("_embedded", {}).get("enheter", []):
            if e.get("konkurs") or e.get("underAvvikling") or e.get("underTvangsavviklingEllerTvangsopplosning"): continue
            nk = e.get("naeringskode1") or {}
            # Bare selskap der koden er hovedaktivitet. Søket treffer også bikoder.
            if not str(nk.get("kode", "")).startswith(kode): continue
            sel.setdefault(e["organisasjonsnummer"], (e, etikett, nk.get("beskrivelse", "")))
        if side + 1 >= r["page"]["totalPages"]: break
        side += 1
seg = lambda a: "A (100+ ansatte)" if a >= 100 else "B (50–99)" if a >= 50 else "C (20–49)"
selskap = []
for nr, (e, et, beskr) in sel.items():
    a = e.get("antallAnsatte") or 0
    adr = e.get("forretningsadresse") or {}
    navn = e["navn"].title() if e["navn"].isupper() else e["navn"]
    selskap.append([navn, nr, et, beskr, a, seg(a), (adr.get("kommune") or "").title(), e.get("hjemmeside", ""),
                    e.get("epostadresse", ""), (e.get("telefon") or e.get("mobil") or "").replace("\xa0", " "),
                    "Ja" if e.get("erIKonsern") else "Nei"])
selskap.sort(key=lambda x: -x[4])

# ---------------------------------------------------------------- Excel
EKSTRA = ["Kontaktperson", "Rolle", "E-post", "Status", "Notat"]
wb = Workbook()
def ark(ws, kol, data, brede):
    ws.append(kol + EKSTRA)
    for r in data: ws.append(r + [""] * len(EKSTRA))
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(vertical="center", wrap_text=True)
        c.fill = PatternFill("solid", fgColor="654B70" if i <= len(kol) else "8A7494")
        ws.column_dimensions[get_column_letter(i)].width = (brede + [22] * len(EKSTRA))[i - 1]
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 32
ws = wb.active; ws.title = "Kommuner"
ark(ws, ["Kommunenr", "Kommune", "Innbyggere", "Segment", "Pris per mnd (12 mnd)", "Årsverdi", "Org.nr", "Nettside", "E-post (Enhetsregisteret)", "Telefon"], kommuner, [11, 24, 12, 20, 16, 14, 13, 32, 34, 16])
for row in ws.iter_rows(min_row=2): row[2].number_format = "# ##0"
ark(wb.create_sheet("Selskap"), ["Selskap", "Org.nr", "Bransje", "Næringskode", "Ansatte", "Segment", "Kommune", "Nettside", "E-post (Enhetsregisteret)", "Telefon", "I konsern"], selskap, [34, 12, 34, 40, 9, 16, 16, 30, 32, 16, 10])
wb.save("kontaktliste.xlsx")
print("kommuner", len(kommuner), "selskap", len(selskap))
